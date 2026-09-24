"""Exporter: saves the approved tailored resume as DOCX and PDF. No LLM here."""

from pathlib import Path

from docx import Document
from fpdf import FPDF

from agents.resume_analyzer.models import TailoredResume
from graph.state import AnalyzerState

OUTPUT_DIR = Path("output")


def _contact_line(resume: TailoredResume) -> str:
    return " | ".join(part for part in (resume.email, resume.phone, resume.location) if part)


def to_docx(resume: TailoredResume, path: Path) -> Path:
    """Write the resume to a Word file."""
    doc = Document()
    doc.add_heading(resume.full_name, level=0)
    doc.add_paragraph(_contact_line(resume))

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(resume.summary)

    doc.add_heading("Skills", level=1)
    doc.add_paragraph(", ".join(resume.skills))

    doc.add_heading("Experience", level=1)
    for job in resume.experience:
        doc.add_paragraph(f"{job.job_title} - {job.company} ({job.dates})").runs[0].bold = True
        for bullet in job.bullets:
            doc.add_paragraph(bullet, style="List Bullet")

    if resume.education:
        doc.add_heading("Education", level=1)
        for item in resume.education:
            doc.add_paragraph(item)

    if resume.certifications:
        doc.add_heading("Certifications", level=1)
        for item in resume.certifications:
            doc.add_paragraph(item)

    doc.save(path)
    return path


# The built-in PDF fonts only support Latin-1, so swap common "smart" characters the LLM likes to use.
_REPLACEMENTS = {"‐": "-", "‑": "-", "–": "-", "—": "-", "‘": "'", "’": "'",
                 "“": '"', "”": '"', "•": "-", "…": "...", " ": " ", " ": " "}


def _pdf_text(text: str) -> str:
    for char, replacement in _REPLACEMENTS.items():
        text = text.replace(char, replacement)
    return text.encode("latin-1", "replace").decode("latin-1")


def to_pdf(resume: TailoredResume, path: Path) -> Path:
    """Write the resume to a PDF file."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(18, 18)

    def heading(text: str) -> None:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 13)
        pdf.multi_cell(0, 7, _pdf_text(text), new_x="LMARGIN", new_y="NEXT")

    def line(text: str, style: str = "") -> None:
        pdf.set_font("Helvetica", style, 10.5)
        pdf.multi_cell(0, 5.5, _pdf_text(text), new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "B", 18)
    pdf.multi_cell(0, 9, _pdf_text(resume.full_name), new_x="LMARGIN", new_y="NEXT")
    line(_contact_line(resume))

    heading("Summary")
    line(resume.summary)

    heading("Skills")
    line(", ".join(resume.skills))

    heading("Experience")
    for job in resume.experience:
        line(f"{job.job_title} - {job.company} ({job.dates})", style="B")
        for bullet in job.bullets:
            line(f"  - {bullet}")
        pdf.ln(1)

    if resume.education:
        heading("Education")
        for item in resume.education:
            line(item)

    if resume.certifications:
        heading("Certifications")
        for item in resume.certifications:
            line(item)

    pdf.output(str(path))
    return path


def exporter_node(state: AnalyzerState) -> dict:
    """LangGraph node: reads `tailored_resume`, writes `export_paths`."""
    OUTPUT_DIR.mkdir(exist_ok=True)
    resume = state["tailored_resume"]
    docx_path = to_docx(resume, OUTPUT_DIR / "tailored_resume.docx")
    pdf_path = to_pdf(resume, OUTPUT_DIR / "tailored_resume.pdf")
    return {"export_paths": {"docx": str(docx_path), "pdf": str(pdf_path)}}
